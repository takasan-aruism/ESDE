"""
ESDE Phase 9: Dictionary Provider
=================================
Psycholinguistic dictionary loader and score provider.

Supported Dictionaries:
  - Brysbaert Concreteness Ratings (40k words)
  - Kuperman Age of Acquisition (30k words)
  - Lancaster Sensorimotor Norms (40k words)
  - Warriner Emotional Valence (14k words)

Design:
  - Singleton pattern for memory efficiency
  - Lazy loading (load on first access)
  - Graceful fallback (Null = -1.0)
  - Lemmatization for better coverage

Null Strategy:
  -1.0 = "Unknown / Out of Vocabulary"
  This is distinct from 0.0 (low score) and signals
  that the word is likely a proper noun, technical term,
  or rare word not in the dictionary.

Spec: Phase 9 W1 Feature Extraction v1.0
"""

import os
import csv
from pathlib import Path
from typing import Dict, Optional, Tuple, List
from threading import Lock

# ==========================================
# Constants
# ==========================================

DICT_PROVIDER_VERSION = "v1.1.0"

# Null value for missing entries
NULL_SCORE = -1.0

# Dictionary types
DICT_CONCRETENESS = "concreteness"
DICT_AOA = "aoa"
DICT_SENSORIMOTOR = "sensorimotor"
DICT_VALENCE = "valence"

# Default dictionary directory
DEFAULT_DICT_DIR = Path(__file__).parent / "dictionaries"

# Dictionary file names (multiple patterns supported)
DICT_FILE_PATTERNS = {
    DICT_CONCRETENESS: [
        "13428_2013_403_MOESM1_ESM.xlsx",  # Original from paper
        "brysbaert_concreteness.xlsx",
        "brysbaert_concreteness.csv",
    ],
    DICT_AOA: [
        "AoA_51715_words.csv",             # Original from paper
        "kuperman_aoa.csv",
    ],
    DICT_SENSORIMOTOR: [
        "Lancaster_sensorimotor_norms_for_39707_words.csv",  # Original
        "lancaster_sensorimotor.csv",
    ],
    DICT_VALENCE: [
        "BRM-emot-submit.csv",             # Original from paper
        "warriner_valence.csv",
    ],
}

# Legacy single-file mapping (for backward compatibility)
DICT_FILES = {
    DICT_CONCRETENESS: "brysbaert_concreteness.csv",
    DICT_AOA: "kuperman_aoa.csv",
    DICT_SENSORIMOTOR: "lancaster_sensorimotor.csv",
    DICT_VALENCE: "warriner_valence.csv",
}

# Score ranges for normalization
SCORE_RANGES = {
    DICT_CONCRETENESS: (1.0, 5.0),    # Brysbaert: 1-5 scale
    DICT_AOA: (0.0, 25.0),            # Kuperman: 0-25 years
    DICT_SENSORIMOTOR: (0.0, 5.0),    # Lancaster: 0-5 scale (max of dimensions)
    DICT_VALENCE: (1.0, 9.0),         # Warriner: 1-9 scale
}


# ==========================================
# Dictionary Provider (Singleton)
# ==========================================

class DictionaryProvider:
    """
    Singleton provider for psycholinguistic dictionaries.
    
    Thread-safe lazy loading with graceful fallback.
    
    Usage:
        provider = DictionaryProvider.get_instance()
        
        # Standard Mode (current)
        score = provider.get_score("sword", "concreteness")
        # Returns 0.0-1.0 if found, -1.0 if not found
        
        # Advanced Mode (future)
        details = provider.get_detailed_scores("sword", "concreteness")
        # Returns {"Conc.M": 4.93, "Conc.SD": 0.26, ...} or None
    """
    
    _instance: Optional["DictionaryProvider"] = None
    _lock = Lock()
    
    def __init__(self, dict_dir: Optional[Path] = None):
        """
        Initialize provider. Use get_instance() instead.
        """
        self.dict_dir = dict_dir or DEFAULT_DICT_DIR
        
        # ==========================================
        # Standard Mode: Primary score only
        # ==========================================
        self._concreteness: Optional[Dict[str, float]] = None
        self._aoa: Optional[Dict[str, float]] = None
        self._sensorimotor: Optional[Dict[str, float]] = None
        self._valence: Optional[Dict[str, float]] = None
        
        # ==========================================
        # Advanced Mode: Full column data (box prepared)
        # ==========================================
        # Structure: {word: {column_name: value, ...}}
        self._concreteness_full: Optional[Dict[str, Dict[str, float]]] = None
        self._aoa_full: Optional[Dict[str, Dict[str, float]]] = None
        self._sensorimotor_full: Optional[Dict[str, Dict[str, float]]] = None
        self._valence_full: Optional[Dict[str, Dict[str, float]]] = None
        
        # Column names stored for reference
        self._column_names: Dict[str, List[str]] = {
            DICT_CONCRETENESS: [],
            DICT_AOA: [],
            DICT_SENSORIMOTOR: [],
            DICT_VALENCE: [],
        }
        
        # Load status
        self._loaded: Dict[str, bool] = {
            DICT_CONCRETENESS: False,
            DICT_AOA: False,
            DICT_SENSORIMOTOR: False,
            DICT_VALENCE: False,
        }
        
        # Stats
        self._stats = {
            "queries": 0,
            "hits": 0,
            "misses": 0,
        }
    
    @classmethod
    def get_instance(cls, dict_dir: Optional[Path] = None) -> "DictionaryProvider":
        """Get singleton instance."""
        if cls._instance is None:
            with cls._lock:
                if cls._instance is None:
                    cls._instance = cls(dict_dir)
        return cls._instance
    
    @classmethod
    def reset_instance(cls) -> None:
        """Reset singleton (for testing)."""
        with cls._lock:
            cls._instance = None
    
    # ==========================================
    # Dictionary Loading
    # ==========================================
    
    def _find_dict_file(self, dict_type: str) -> Optional[Path]:
        """Find dictionary file from multiple patterns."""
        patterns = DICT_FILE_PATTERNS.get(dict_type, [])
        
        # Also check legacy single filename
        legacy = DICT_FILES.get(dict_type)
        if legacy and legacy not in patterns:
            patterns.append(legacy)
        
        for pattern in patterns:
            path = self.dict_dir / pattern
            if path.exists():
                return path
        
        return None
    
    def _parse_eu_number(self, value: str) -> Optional[float]:
        """Parse EU-style number (1,46 → 1.46)."""
        if not value:
            return None
        try:
            # Replace comma with dot for EU format
            cleaned = value.replace(',', '.')
            return float(cleaned)
        except ValueError:
            return None
    
    def _load_xlsx(self, path: Path, word_col: str, value_col: str) -> Dict[str, float]:
        """Load dictionary from XLSX file."""
        cache: Dict[str, float] = {}
        
        try:
            import openpyxl
        except ImportError:
            print(f"[DictProvider] Warning: openpyxl not installed. Cannot read {path}")
            print(f"[DictProvider] Install with: pip install openpyxl")
            return cache
        
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            
            # Find column indices from header row
            header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            
            word_idx = None
            value_idx = None
            for i, col_name in enumerate(header):
                if col_name == word_col:
                    word_idx = i
                elif col_name == value_col:
                    value_idx = i
            
            if word_idx is None or value_idx is None:
                print(f"[DictProvider] Warning: Columns not found in {path}")
                print(f"[DictProvider] Looking for: {word_col}, {value_col}")
                print(f"[DictProvider] Found: {header}")
                return cache
            
            # Read data rows
            for row in ws.iter_rows(min_row=2, values_only=True):
                if len(row) <= max(word_idx, value_idx):
                    continue
                
                word = row[word_idx]
                value = row[value_idx]
                
                if not word:
                    continue
                
                word = str(word).strip().lower()
                
                # Handle EU number format or regular number
                if isinstance(value, (int, float)):
                    cache[word] = float(value)
                elif isinstance(value, str):
                    parsed = self._parse_eu_number(value)
                    if parsed is not None:
                        cache[word] = parsed
            
            wb.close()
            print(f"[DictProvider] Loaded {len(cache)} entries from {path.name}")
            
        except Exception as e:
            print(f"[DictProvider] Error loading {path}: {e}")
        
        return cache
    
    def _load_concreteness(self) -> Dict[str, float]:
        """
        Load Brysbaert Concreteness Ratings.
        
        Supports:
          - XLSX format (13428_2013_403_MOESM1_ESM.xlsx)
          - CSV format (brysbaert_concreteness.csv)
          
        Standard: Conc.M for primary score
        Advanced: Conc.M, Conc.SD stored in _concreteness_full
        """
        path = self._find_dict_file(DICT_CONCRETENESS)
        cache: Dict[str, float] = {}
        cache_full: Dict[str, Dict[str, float]] = {}
        
        if not path:
            print(f"[DictProvider] Warning: Concreteness dictionary not found")
            return cache
        
        # XLSX format (primary value only for now)
        if path.suffix.lower() == '.xlsx':
            cache = self._load_xlsx(path, 'Word', 'Conc.M')
            # Note: Full columns not loaded for XLSX yet
            # Future: extend _load_xlsx or convert to CSV
            return cache
        
        # CSV format - load all columns
        primary_col = 'Conc.M'
        numeric_cols = ['Conc.M', 'Conc.SD']
        
        try:
            with open(path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                
                # Store column names for reference
                self._column_names[DICT_CONCRETENESS] = list(reader.fieldnames or [])
                
                for row in reader:
                    word = row.get('Word', '').strip().lower()
                    if not word:
                        continue
                    
                    # Standard: Conc.M
                    score_str = row.get(primary_col, '')
                    if score_str:
                        parsed = self._parse_eu_number(score_str)
                        if parsed is not None:
                            cache[word] = parsed
                    
                    # Advanced: store Conc.M and Conc.SD
                    full_entry: Dict[str, float] = {}
                    for col in numeric_cols:
                        val_str = row.get(col, '')
                        if val_str:
                            parsed = self._parse_eu_number(val_str)
                            if parsed is not None:
                                full_entry[col] = parsed
                    if full_entry:
                        cache_full[word] = full_entry
            
            self._concreteness_full = cache_full
            print(f"[DictProvider] Loaded {len(cache)} concreteness entries")
        except Exception as e:
            print(f"[DictProvider] Error loading concreteness: {e}")
        
        return cache
    
    def _load_aoa(self) -> Dict[str, float]:
        """
        Load Kuperman Age of Acquisition norms.
        
        Supports multiple CSV formats:
          - AoA_51715_words.csv (Word, Rating.Mean)
          - kuperman_aoa.csv
        """
        path = self._find_dict_file(DICT_AOA)
        cache: Dict[str, float] = {}
        
        if not path:
            print(f"[DictProvider] Warning: AoA dictionary not found")
            return cache
        
        try:
            with open(path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                
                # Detect column name (Rating.Mean or AoA_Kup_lem or similar)
                fieldnames = reader.fieldnames or []
                score_col = None
                for col in ['Rating.Mean', 'AoA_Kup_lem', 'AoA', 'Rating']:
                    if col in fieldnames:
                        score_col = col
                        break
                
                if not score_col:
                    print(f"[DictProvider] Warning: AoA score column not found")
                    print(f"[DictProvider] Available columns: {fieldnames}")
                    return cache
                
                for row in reader:
                    word = row.get('Word', '').strip().lower()
                    score_str = row.get(score_col, '')
                    if word and score_str:
                        try:
                            cache[word] = float(score_str)
                        except ValueError:
                            pass
            print(f"[DictProvider] Loaded {len(cache)} AoA entries")
        except Exception as e:
            print(f"[DictProvider] Error loading AoA: {e}")
        
        return cache
    
    def _load_sensorimotor(self) -> Dict[str, float]:
        """
        Load Lancaster Sensorimotor Norms.
        
        File: Lancaster_sensorimotor_norms_for_39707_words.csv
        
        Columns: Word, Auditory.mean, Gustatory.mean, Haptic.mean,
                 Olfactory.mean, Visual.mean, ...
                 
        Standard: max(perceptual) for primary score
        Advanced: all columns stored in _sensorimotor_full
        """
        path = self._find_dict_file(DICT_SENSORIMOTOR)
        cache: Dict[str, float] = {}
        cache_full: Dict[str, Dict[str, float]] = {}
        
        if not path:
            print(f"[DictProvider] Warning: Sensorimotor dictionary not found")
            return cache
        
        # Columns for primary score (Standard mode)
        perceptual_cols = [
            'Auditory.mean', 'Gustatory.mean', 'Haptic.mean',
            'Olfactory.mean', 'Visual.mean'
        ]
        
        # All numeric columns to preserve (Advanced mode)
        numeric_cols = [
            'Auditory.mean', 'Gustatory.mean', 'Haptic.mean',
            'Interoceptive.mean', 'Olfactory.mean', 'Visual.mean',
            'Foot_leg.mean', 'Hand_arm.mean', 'Head.mean',
            'Mouth.mean', 'Torso.mean',
            'Auditory.SD', 'Gustatory.SD', 'Haptic.SD',
            'Interoceptive.SD', 'Olfactory.SD', 'Visual.SD',
            'Foot_leg.SD', 'Hand_arm.SD', 'Head.SD',
            'Mouth.SD', 'Torso.SD',
        ]
        
        try:
            with open(path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                
                # Store column names for reference
                self._column_names[DICT_SENSORIMOTOR] = list(reader.fieldnames or [])
                
                for row in reader:
                    word = row.get('Word', '').strip().lower()
                    if not word:
                        continue
                    
                    # Standard: max of perceptual dimensions
                    max_score = 0.0
                    for col in perceptual_cols:
                        score_str = row.get(col, '')
                        if score_str:
                            try:
                                max_score = max(max_score, float(score_str))
                            except ValueError:
                                pass
                    cache[word] = max_score
                    
                    # Advanced: store all numeric columns
                    full_entry: Dict[str, float] = {}
                    for col in numeric_cols:
                        val_str = row.get(col, '')
                        if val_str:
                            try:
                                full_entry[col] = float(val_str)
                            except ValueError:
                                pass
                    if full_entry:
                        cache_full[word] = full_entry
            
            self._sensorimotor_full = cache_full
            print(f"[DictProvider] Loaded {len(cache)} sensorimotor entries ({len(numeric_cols)} columns preserved)")
        except Exception as e:
            print(f"[DictProvider] Error loading sensorimotor: {e}")
        
        return cache
    
    def _load_valence(self) -> Dict[str, float]:
        """
        Load Warriner Emotional Valence norms.
        
        File: BRM-emot-submit.csv
        
        Format: (index),Word,V.Mean.Sum,V.SD.Sum,...
        Note: First column is row index (unnamed), skip it.
        
        Standard: V.Mean.Sum for primary score
        Advanced: V/A/D all stored in _valence_full
        """
        path = self._find_dict_file(DICT_VALENCE)
        cache: Dict[str, float] = {}
        cache_full: Dict[str, Dict[str, float]] = {}
        
        if not path:
            print(f"[DictProvider] Warning: Valence dictionary not found")
            return cache
        
        # Primary column (Standard mode)
        primary_col = 'V.Mean.Sum'
        
        # All VAD columns to preserve (Advanced mode)
        # V = Valence (positive/negative)
        # A = Arousal (excited/calm)
        # D = Dominance (strong/weak)
        vad_cols = [
            'V.Mean.Sum', 'V.SD.Sum',
            'A.Mean.Sum', 'A.SD.Sum',
            'D.Mean.Sum', 'D.SD.Sum',
        ]
        
        try:
            with open(path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                
                # Store column names for reference
                self._column_names[DICT_VALENCE] = list(reader.fieldnames or [])
                
                for row in reader:
                    word = row.get('Word', '').strip().lower()
                    if not word:
                        continue
                    
                    # Standard: V.Mean.Sum only
                    score_str = row.get(primary_col, '')
                    if score_str:
                        try:
                            cache[word] = float(score_str)
                        except ValueError:
                            pass
                    
                    # Advanced: store V/A/D columns
                    full_entry: Dict[str, float] = {}
                    for col in vad_cols:
                        val_str = row.get(col, '')
                        if val_str:
                            try:
                                full_entry[col] = float(val_str)
                            except ValueError:
                                pass
                    if full_entry:
                        cache_full[word] = full_entry
            
            self._valence_full = cache_full
            print(f"[DictProvider] Loaded {len(cache)} valence entries (V/A/D preserved)")
        except Exception as e:
            print(f"[DictProvider] Error loading valence: {e}")
        
        return cache
    
    def _ensure_loaded(self, dict_type: str) -> None:
        """Ensure dictionary is loaded (lazy loading)."""
        if self._loaded.get(dict_type):
            return
        
        with self._lock:
            if self._loaded.get(dict_type):
                return
            
            if dict_type == DICT_CONCRETENESS:
                self._concreteness = self._load_concreteness()
            elif dict_type == DICT_AOA:
                self._aoa = self._load_aoa()
            elif dict_type == DICT_SENSORIMOTOR:
                self._sensorimotor = self._load_sensorimotor()
            elif dict_type == DICT_VALENCE:
                self._valence = self._load_valence()
            
            self._loaded[dict_type] = True
    
    def _get_cache(self, dict_type: str) -> Optional[Dict[str, float]]:
        """Get dictionary cache."""
        self._ensure_loaded(dict_type)
        
        if dict_type == DICT_CONCRETENESS:
            return self._concreteness
        elif dict_type == DICT_AOA:
            return self._aoa
        elif dict_type == DICT_SENSORIMOTOR:
            return self._sensorimotor
        elif dict_type == DICT_VALENCE:
            return self._valence
        
        return None
    
    # ==========================================
    # Public API
    # ==========================================
    
    def get_raw_score(self, word: str, dict_type: str) -> Optional[float]:
        """
        Get raw (unnormalized) score for a word.
        
        Args:
            word: Word to look up
            dict_type: Dictionary type (concreteness, aoa, sensorimotor, valence)
            
        Returns:
            Raw score if found, None if not found
        """
        cache = self._get_cache(dict_type)
        if cache is None:
            return None
        
        # Try exact match first
        normalized = word.lower().strip()
        return cache.get(normalized)
    
    def get_score(
        self,
        word: str,
        dict_type: str,
        lemma: Optional[str] = None,
    ) -> float:
        """
        Get normalized score for a word.
        
        Normalization: (raw - min) / (max - min) → [0.0, 1.0]
        
        Args:
            word: Word to look up
            dict_type: Dictionary type
            lemma: Optional lemma form to try if word not found
            
        Returns:
            Normalized score [0.0, 1.0] if found, -1.0 if not found
        """
        self._stats["queries"] += 1
        
        # Try word first
        raw = self.get_raw_score(word, dict_type)
        
        # Try lemma if word not found
        if raw is None and lemma and lemma != word.lower():
            raw = self.get_raw_score(lemma, dict_type)
        
        if raw is None:
            self._stats["misses"] += 1
            return NULL_SCORE
        
        self._stats["hits"] += 1
        
        # Normalize to [0.0, 1.0]
        min_val, max_val = SCORE_RANGES.get(dict_type, (0.0, 1.0))
        normalized = (raw - min_val) / (max_val - min_val)
        
        # Clamp to [0.0, 1.0]
        return max(0.0, min(1.0, normalized))
    
    def get_all_scores(
        self,
        word: str,
        lemma: Optional[str] = None,
    ) -> Dict[str, float]:
        """
        Get all dictionary scores for a word.
        
        Args:
            word: Word to look up
            lemma: Optional lemma form
            
        Returns:
            Dict with all scores (each -1.0 if not found)
        """
        return {
            DICT_CONCRETENESS: self.get_score(word, DICT_CONCRETENESS, lemma),
            DICT_AOA: self.get_score(word, DICT_AOA, lemma),
            DICT_SENSORIMOTOR: self.get_score(word, DICT_SENSORIMOTOR, lemma),
            DICT_VALENCE: self.get_score(word, DICT_VALENCE, lemma),
        }
    
    def get_stats(self) -> Dict[str, any]:
        """Get usage statistics."""
        total = self._stats["queries"]
        hit_rate = self._stats["hits"] / total if total > 0 else 0.0
        
        return {
            "version": DICT_PROVIDER_VERSION,
            "queries": total,
            "hits": self._stats["hits"],
            "misses": self._stats["misses"],
            "hit_rate": hit_rate,
            "loaded": dict(self._loaded),
            "sizes": {
                DICT_CONCRETENESS: len(self._concreteness) if self._concreteness else 0,
                DICT_AOA: len(self._aoa) if self._aoa else 0,
                DICT_SENSORIMOTOR: len(self._sensorimotor) if self._sensorimotor else 0,
                DICT_VALENCE: len(self._valence) if self._valence else 0,
            }
        }
    
    # ==========================================
    # Advanced API (Future Extension)
    # ==========================================
    
    def get_detailed_scores(
        self,
        word: str,
        dict_type: str,
        lemma: Optional[str] = None,
    ) -> Optional[Dict[str, float]]:
        """
        Get all columns for a word from a dictionary.
        
        [ADVANCED MODE - For future extension]
        
        Args:
            word: Word to look up
            dict_type: Dictionary type
            lemma: Optional lemma form
            
        Returns:
            Dict of {column_name: raw_value} or None if not found
        """
        self._ensure_loaded(dict_type)
        
        normalized = word.lower().strip()
        
        # Get from full cache
        full_cache = None
        if dict_type == DICT_CONCRETENESS:
            full_cache = self._concreteness_full
        elif dict_type == DICT_AOA:
            full_cache = self._aoa_full
        elif dict_type == DICT_SENSORIMOTOR:
            full_cache = self._sensorimotor_full
        elif dict_type == DICT_VALENCE:
            full_cache = self._valence_full
        
        if full_cache is None:
            return None
        
        # Try word first
        result = full_cache.get(normalized)
        
        # Try lemma if not found
        if result is None and lemma and lemma.lower() != normalized:
            result = full_cache.get(lemma.lower())
        
        return result
    
    def get_column_names(self, dict_type: str) -> List[str]:
        """
        Get available column names for a dictionary.
        
        [ADVANCED MODE - For future extension]
        
        Returns:
            List of column names, or empty list if not loaded
        """
        self._ensure_loaded(dict_type)
        return self._column_names.get(dict_type, [])
    
    def is_loaded(self, dict_type: str) -> bool:
        """Check if dictionary is loaded."""
        return self._loaded.get(dict_type, False)
    
    def preload_all(self) -> None:
        """Preload all dictionaries."""
        for dict_type in DICT_FILES.keys():
            self._ensure_loaded(dict_type)


# ==========================================
# Convenience Functions
# ==========================================

def get_concreteness(word: str, lemma: Optional[str] = None) -> float:
    """Get concreteness score."""
    return DictionaryProvider.get_instance().get_score(word, DICT_CONCRETENESS, lemma)

def get_aoa(word: str, lemma: Optional[str] = None) -> float:
    """Get age of acquisition score."""
    return DictionaryProvider.get_instance().get_score(word, DICT_AOA, lemma)

def get_sensorimotor(word: str, lemma: Optional[str] = None) -> float:
    """Get sensorimotor intensity score."""
    return DictionaryProvider.get_instance().get_score(word, DICT_SENSORIMOTOR, lemma)

def get_valence(word: str, lemma: Optional[str] = None) -> float:
    """Get emotional valence score."""
    return DictionaryProvider.get_instance().get_score(word, DICT_VALENCE, lemma)


# ==========================================
# Test
# ==========================================

if __name__ == "__main__":
    print("Dictionary Provider Test")
    print("=" * 60)
    
    provider = DictionaryProvider.get_instance()
    
    # Test words
    test_words = [
        ("sword", None),
        ("honor", None),
        ("kill", "kill"),
        ("killed", "kill"),
        ("beautiful", None),
        ("Nobunaga", None),  # Proper noun, likely missing
        ("asdfghjkl", None),  # Nonsense, definitely missing
    ]
    
    print("\n[Test] Score lookups:")
    for word, lemma in test_words:
        scores = provider.get_all_scores(word, lemma)
        print(f"\n  '{word}' (lemma={lemma}):")
        for dict_type, score in scores.items():
            status = "NULL" if score == NULL_SCORE else f"{score:.3f}"
            print(f"    {dict_type}: {status}")
    
    print("\n[Stats]")
    stats = provider.get_stats()
    for k, v in stats.items():
        print(f"  {k}: {v}")
    
    print("\n" + "=" * 60)
    print("Test complete!")
